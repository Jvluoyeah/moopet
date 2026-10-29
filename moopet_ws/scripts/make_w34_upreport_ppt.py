#!/usr/bin/env python3
"""生成 W3+W4 合并工作汇报 PPT（李文扬 · 架构/集成，向上级汇报版）。

背景：上周领导有事未听 W3 汇报，本周 W3+W4 一起汇报。
用法：
  python scripts/make_w34_upreport_ppt.py   # 输出 W3-W4工作汇报_李文扬.pptx
依赖：python-pptx（ppt-master/venv 已装）
"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'W3-W4工作汇报_李文扬.pptx')

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
CYAN = RGBColor(0x3B, 0xA7, 0xA5)
ORANGE = RGBColor(0xE8, 0x7A, 0x2E)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC0, 0x39, 0x2B)
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)
GRAY = RGBColor(0x55, 0x5F, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = '微软雅黑'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 11


def _set_ea(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


def set_run(run, size, bold=False, color=NAVY, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    _set_ea(run)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, it in enumerate(lines):
        text, size, bold, color = it[0], it[1], it[2], it[3]
        sa = it[4] if len(it) > 4 else space_after
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sa)
        r = p.add_run()
        r.text = text
        set_run(r, size, bold, color)
    return tb


def add_rect(slide, x, y, w, h, fill, line=None, text='', tsize=14, tcolor=WHITE,
             bold=True, rounded=False, line_w=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_run(r, tsize, bold, tcolor)
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=BLUE, w=2.25):
    conn = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(w)
    return conn


def title_bar(slide, title, sub=''):
    add_rect(slide, 0, 0, 13.333, 0.95, NAVY, text=title, tsize=23)
    add_rect(slide, 0, 0.95, 13.333, 0.06, ORANGE)
    if sub:
        add_text(slide, 0.5, 1.12, 12.3, 0.4, [(sub, 12.5, False, GRAY)])


def footer(slide, idx):
    add_text(slide, 11.4, 7.1, 1.7, 0.3,
             [(f'{idx} / {TOTAL}', 10, False, GRAY)], align=PP_ALIGN.RIGHT)


def add_table(slide, rows, x, y, w, h, col_widths, header_size=11.5, body_size=10.5,
              row_h=None):
    ts = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                Inches(w), Inches(h))
    tb = ts.table
    for ci, wd in enumerate(col_widths):
        tb.columns[ci].width = Inches(wd)
    for ri, row in enumerate(rows):
        if row_h and ri == 0:
            tb.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = tb.cell(ri, ci)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = val
            set_run(r, header_size if ri == 0 else body_size, ri == 0,
                    WHITE if ri == 0 else NAVY)
    return tb


# ============================================================ 1 封面
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, 13.333, 7.5, NAVY)
add_rect(s, 0, 4.75, 13.333, 0.08, ORANGE)
add_text(s, 1.0, 1.2, 11.3, 1.1, [('W3+W4 工作汇报（8.17–8.30）', 38, True, WHITE)])
add_text(s, 1.0, 2.35, 11.3, 0.8,
         [('答题闭环 → 情感决策完稿 · 双端联调 · 成长日记', 21, True, CYAN)])
add_text(s, 1.0, 3.5, 11.3, 1.0, [
    ('汇报人：李文扬（架构 / 集成）', 16, False, WHITE),
    ('长夜月 · 蓝思杯赛道一：灵伴管家（智能体应用）', 14, False, WHITE),
    ('对齐排期 W3+W4 · 李文扬条目（上周 W3 因故未汇报，本周合并汇报）', 13, False, WHITE),
])
add_text(s, 1.0, 5.3, 11.3, 1.2, [
    ('双周结论：', 15, True, ORANGE),
    ('W3 答题闭环 + W4 决策引擎完稿均按排期完成，两周连起来就是一条完整成长链：', 15, False, WHITE),
    ('感知互动 → 加点升级 → 答题锻炼 → 性格养成 → 情感陪伴 → 语音说出来（47/47 单测）。', 15, False, WHITE),
])
add_text(s, 1.0, 6.85, 11.3, 0.4,
         [('代码：w4/moopet_ws（周周快照，W1-W3 成果全部保留）', 12, False, WHITE)])

# ============================================================ 2 双周结论（一页看懂）
s = prs.slides.add_slide(BLANK)
title_bar(s, '双周结论（一页看懂）')
# W3 卡
add_rect(s, 0.6, 1.6, 5.9, 3.0, LIGHT, rounded=True)
add_rect(s, 0.6, 1.6, 5.9, 0.62, BLUE, rounded=True, text='W3（8.17–8.23）答题 + 基础闭环', tsize=14)
add_text(s, 0.85, 2.4, 5.45, 2.1, [
    ('· quiz_node 集成：出题/判题/连击/难度降档', 12, False, NAVY),
    ('· 行为决策引擎骨架（BT 目录 + 空节点）', 12, False, NAVY),
    ('· 挂到决策入口：作答事件 → 决策反应', 12, False, NAVY),
    ('· 单测 28/28；16 项接口；/dog/combo 桥接', 12, False, NAVY),
    ('· demo 终态 LV.23 / 1300 经验', 12, False, NAVY),
])
# W4 卡
add_rect(s, 6.85, 1.6, 5.9, 3.0, LIGHT, rounded=True)
add_rect(s, 6.85, 1.6, 5.9, 0.62, GREEN, rounded=True, text='W4（8.24–8.30）情感决策 + 双端联调', tsize=14)
add_text(s, 7.1, 2.4, 5.45, 2.1, [
    ('· 决策引擎完稿：性格演化（四型）+ 情感账户', 12, False, NAVY),
    ('· 主动关怀（长沉默/深夜）+ 早安晚安/纪念日', 12, False, NAVY),
    ('· 成长日记：按天统计 → 第一人称日记落库', 12, False, NAVY),
    ('· 背包 TTS 节点：决策输出 → 语音+动作 双端联调', 12, False, NAVY),
    ('· 单测累计 47/47（新增 19）', 12, False, NAVY),
])
# 底部结论条
add_rect(s, 0.6, 4.9, 12.1, 1.75, NAVY, rounded=True)
add_text(s, 0.95, 5.12, 11.5, 1.4, [
    ('两周合起来看：', 14.5, True, ORANGE),
    ('W3 让狗"会做题"，W4 让狗"有性格、懂感情、会主动开口"——排期 9.18 前的软件主线', 13, False, WHITE),
    ('（感知→成长→答题→情感→语音）已全部打通，剩下的是 W5-W6 的火花计划与稳定性收尾。', 13, False, WHITE),
])
footer(s, 2)

# ============================================================ 3 排期对照
s = prs.slides.add_slide(BLANK)
title_bar(s, '排期对照（两周 · 李文扬条目）')
rows = [
    ('周次', '排期任务', '完成情况', '状态'),
    ('W3', '集成 quiz_node 并挂到决策引擎入口', '答题节点接入决策入口；连击/难度/桥接齐备', '✅'),
    ('W3', '搭建行为决策引擎骨架（BT 目录 + 空节点）', 'core/conditions/actions/trees 四件套 + 入口节点', '✅'),
    ('W4', '完成行为决策引擎（性格/主动关怀/纪念日）', '双轴四型性格演化 + 主动关怀 + 仪式/纪念日（真实化 W3 空节点）', '✅'),
    ('W4', '情感账户', '0~100 账户：互动+ / 闲置- / 冷淡·平衡·依赖三档 / 和好机制', '✅'),
    ('W4', '双端联调（/dog/action + TTS）', '新增背包 TTS 节点，决策输出 → 动作 + 语音全链路', '✅'),
    ('附加', '成长日记（排期归日记线，先落库）', '按天统计互动 → 第一人称日记 → diary 表（一天一条）', '✅'),
]
add_table(s, rows, 0.6, 1.55, 12.1, 4.3, [0.9, 4.2, 5.6, 1.4])
add_text(s, 0.6, 6.2, 12.1, 0.6, [
    ('说明：两周 5 项排期任务全部完成，无延期；附加完成日记落库，为牛浩然 W4 日记 API 直接铺路。', 12.5, False, GRAY),
])
footer(s, 3)

# ============================================================ 4 W3 浓缩回顾
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W3 浓缩回顾（上周未汇报，这里一页带过）')
add_table(s, [
    ('W3 交付', '要点', '与 W4 的衔接'),
    ('quiz_node 答题闭环', '等级解锁 3/8/15/20；连击×5→+50、×10→+100；连错 3 题难度降档', '作答事件持续喂给决策引擎 → W4 性格"表达度"的来源'),
    ('决策引擎骨架', 'BT 目录四件套 + 入口节点（事件路由 → 反应）', 'W4 只做"填空"：真实化空条件/空动作，树结构不动'),
    ('16 项接口 + 28/28 单测', 'DecisionEventMsg 入口；题库/答题/决策可单测', 'W4 无新增接口，直接在旧接口上实现行为'),
    ('demo 终态', 'LV.23 / 1300 经验 / 本周 380/2000 / 4 项成就', '成长数据已落库，W4 情感层叠加其上'),
], 0.6, 1.55, 12.1, 3.6, [2.6, 5.4, 4.1])
add_rect(s, 0.6, 5.45, 12.1, 1.2, LIGHT, rounded=True)
add_text(s, 0.9, 5.62, 11.6, 0.95, [
    ('一句话：W3 的骨架是"留好了插口"的骨架——W4 的性格/关怀/纪念日就是插上去的模块，', 12.5, False, GRAY),
    ('两周代码同构演进（w4 快照承接 w3），没有返工。', 12.5, False, GRAY),
])
footer(s, 4)

# ============================================================ 5 W4 成果① 性格演化
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W4 成果① 性格演化：双轴 → 四型（对齐狗2.md 独特性格塑造）')
# 左：2×2 四象限
cx, cy = 3.6, 4.05          # 象限中心
add_rect(s, 0.7, 1.7, 5.8, 4.7, LIGHT, rounded=True)
add_rect(s, cx, 1.95, 0.03, 4.2, GRAY)
add_rect(s, 0.95, cy, 5.3, 0.03, GRAY)
add_text(s, 0.85, 1.78, 5.5, 0.3, [('依恋度 高 ↑（陪伴类：摸头/握手/散步/指引）', 10.5, True, BLUE)])
add_text(s, 3.4, 6.02, 3.1, 0.3, [('表达度 高 →（说话类：答题/呼唤）', 10.5, True, ORANGE)])
add_rect(s, 1.15, 2.2, 2.25, 1.55, CYAN, rounded=True, text='安静型\n高依恋 · 低表达\n"默默陪着主人"', tsize=11)
add_rect(s, 3.75, 2.2, 2.25, 1.55, ORANGE, rounded=True, text='粘人型\n高依恋 · 高表达\n"走到哪跟到哪"', tsize=11)
add_rect(s, 1.15, 4.35, 2.25, 1.55, BLUE, rounded=True, text='独立型\n低依恋 · 低表达\n"自己玩也很开心"', tsize=11)
add_rect(s, 3.75, 4.35, 2.25, 1.55, GREEN, rounded=True, text='话痨型\n低依恋 · 高表达\n"有说不完的话"', tsize=11)
# 右：机制与演示
add_rect(s, 6.95, 1.7, 5.8, 2.5, LIGHT, rounded=True)
add_text(s, 7.2, 1.85, 5.35, 2.25, [
    ('机制（计数器 + 衰减）：', 13.5, True, NAVY),
    ('· 命中事件：所在轴 +6，双轴自然衰减 -1', 11.5, False, GRAY),
    ('· 每闲置一天：双轴 -3（性格慢慢淡化）', 11.5, False, GRAY),
    ('· 轴分界 55：双高=粘人 / 高表达=话痨', 11.5, False, GRAY),
    ('  高依恋=安静 / 双低=独立（初始独立型）', 11.5, False, GRAY),
])
add_rect(s, 6.95, 4.45, 5.8, 1.95, GREEN, rounded=True)
add_text(s, 7.2, 4.6, 5.35, 1.7, [
    ('demo 实测演化路径：', 13, True, WHITE),
    ('独立型 → 摸头×4 → 安静型 → 呼唤+答题 → 粘人型', 11.5, False, WHITE),
    ('性格变化时狗狗主动播报："好像…我更黏主人啦，请多指教！"', 11.5, False, WHITE),
])
footer(s, 5)

# ============================================================ 6 W4 成果② 情感账户 + 主动关怀
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W4 成果② 情感账户 + 主动关怀（对齐狗2.md 情感连接层）')
# 账户刻度条
add_text(s, 0.7, 1.5, 6.0, 0.4, [('情感账户 0~100（投入越多回应越丰富）', 13.5, True, NAVY)])
add_rect(s, 0.7, 2.0, 2.2, 0.55, RED, text='0–30 冷淡', tsize=11.5)
add_rect(s, 2.95, 2.0, 2.2, 0.55, BLUE, text='30–70 平衡', tsize=11.5)
add_rect(s, 5.2, 2.0, 2.2, 0.55, GREEN, text='70–100 依赖', tsize=11.5)
add_text(s, 0.7, 2.75, 6.7, 1.7, [
    ('· 每次正向互动 +2（摸头/握手/散步/答对…），封顶 100', 11.5, False, GRAY),
    ('· 每闲置一天 -5；双轴性格同步淡化 -3/天', 11.5, False, GRAY),
    ('· 冷淡期：主动关怀整体抑制（"长期忽视减少主动互动"）', 11.5, False, GRAY),
    ('· 和好机制：回来多互动，账户逐步回温（demo：6 次摸头 20→32）', 11.5, False, GRAY),
])
# 主动关怀
add_text(s, 7.2, 1.5, 5.5, 0.4, [('主动关怀（长沉默 ≥2h 触发）', 13.5, True, NAVY)])
add_rect(s, 7.2, 2.0, 5.5, 2.45, LIGHT, rounded=True)
add_text(s, 7.45, 2.15, 5.0, 2.2, [
    ('· 平衡期："主人好久没理我了…有点想你"', 11.5, False, GRAY),
    ('· 依赖期："主人你去哪啦，我等你好久啦！"', 11.5, False, GRAY),
    ('· 深夜 23-6 点：轻声"夜深了，主人早点休息哦"', 11.5, False, GRAY),
    ('· 冷淡期：沉默不搭话（设计行为，非缺陷）', 11.5, False, GRAY),
    ('· 触发入口：on_tick 周期检查（节点 60s 定时器）', 11.5, False, GRAY),
])
add_rect(s, 0.6, 4.85, 12.1, 1.75, NAVY, rounded=True)
add_text(s, 0.95, 5.05, 11.5, 1.4, [
    ('demo 实测（冷落 → 和好）：', 13.5, True, ORANGE),
    ('2 天不互动 账户 64→54 → 冷落区 20（不再主动）→ 回来摸头×6 → 32 回到平衡期。', 12.5, False, WHITE),
    ('全程数值可复现、可单测（19 项 W4 单测覆盖账户增益/封顶/衰减/档位/抑制）。', 12.5, False, WHITE),
])
footer(s, 6)

# ============================================================ 7 W4 成果③ 仪式/纪念日 + 成长日记
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W4 成果③ 陪伴仪式 + 纪念日 + 成长日记')
items = [
    ('早安 / 晚安', '6-9 点 / 21-23 点窗口，每天一次不重复；按当前性格说不同的话（亮点）'),
    ('等级里程碑', 'LV.10/20/…/100 升级即庆祝："我升到 LV.10 啦！撒花庆祝！"'),
    ('自定义纪念日', "主人可登记（如'认识主人满月'），当天主动提起，每日一次"),
    ('成长日记', '按天统计互动 → 生成第一人称日记 → diary 表落库（一天一条，跨日自动写）'),
]
for i, (h, d) in enumerate(items):
    y = 1.55 + i * 0.78
    add_rect(s, 0.6, y, 2.3, 0.62, NAVY, rounded=True, text=h, tsize=12.5)
    add_rect(s, 3.05, y, 9.65, 0.62, LIGHT, rounded=True)
    add_text(s, 3.25, y + 0.05, 9.3, 0.55, [(d, 11.5, False, NAVY)])
# 性格化话术对比（亮点框）
add_rect(s, 0.6, 4.85, 5.9, 1.9, CYAN, rounded=True)
add_text(s, 0.85, 5.0, 5.45, 1.6, [
    ('同一句早安，性格不同说法不同：', 12.5, True, WHITE),
    ('粘人型："主人早上好呀！我在这儿等你好久啦！"', 11, False, WHITE),
    ('独立型："早啊主人，我去晒会儿太阳。"', 11, False, WHITE),
    ('→ 狗2.md"独特性格塑造"落到每一句话上', 10.5, False, WHITE),
])
add_rect(s, 6.85, 4.85, 5.9, 1.9, LIGHT, rounded=True)
add_text(s, 7.1, 5.0, 5.45, 1.6, [
    ('日记示例（demo 自动生成）：', 12.5, True, NAVY),
    ('"2025-08-28。今天我 LV.10，心情平静，是粘人型的小狗。', 10.5, False, GRAY),
    ('主人叫了我 1 次，我都有好好回应；迎来了升级里程碑 1 次。', 10.5, False, GRAY),
    ('我们的情感账户存到了 30，感觉我们越来越有默契啦。"', 10.5, False, GRAY),
])
footer(s, 7)

# ============================================================ 8 W4 成果④ 双端联调 TTS
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W4 成果④ 双端联调：决策输出 → 动作 + 语音（/dog/action + TTS）')
# 链路图
nodes = [
    (0.6, '感知 / 答题 /\n定时主动检查', BLUE),
    (3.35, '行为决策引擎\n（狗侧）反应：\ntts_text + action', GREEN),
    (6.35, '背包 TTS 节点\n/backpack/tts_speak\n（TtsSpeak 动作）', ORANGE),
    (9.35, '语音播报 🔊\n进度反馈 → 时长结果', CYAN),
    (9.35, '/dog/action\n动作执行 + 加点', BLUE),
]
y1 = 2.2
for x, txt, color in nodes[:4]:
    add_rect(s, x, y1, 2.5, 1.5, color, rounded=True, text=txt, tsize=11)
add_arrow(s, 3.1, y1 + 0.75, 3.35, y1 + 0.75, w=2.5)
add_arrow(s, 5.85, y1 + 0.75, 6.35, y1 + 0.75, w=2.5)
add_arrow(s, 8.85, y1 + 0.75, 9.35, y1 + 0.75, w=2.5)
add_rect(s, 6.35, y1 + 2.3, 2.5, 1.1, BLUE, rounded=True, text=nodes[4][1], tsize=11)
add_arrow(s, 4.6, y1 + 1.5, 6.35, y1 + 2.85, w=2.0)
# 要点
add_rect(s, 9.35, 4.5, 3.35, 2.2, LIGHT, rounded=True)
add_text(s, 9.55, 4.65, 3.0, 1.95, [
    ('W3 只留了客户端：', 11.5, True, NAVY),
    ('"TTS 未就绪记日志"', 11, False, GRAY),
    ('W4 补上服务端即通：', 11.5, True, NAVY),
    ('接口 0 新增，demo 模式', 11, False, GRAY),
    ('打印 🔊 播报内容', 11, False, GRAY),
])
add_rect(s, 0.6, 6.15, 8.5, 0.95, LIGHT, rounded=True)
add_text(s, 0.85, 6.3, 8.1, 0.7, [
    ('联调口径：goal(text/rate/pitch) → feedback(progress) → result(success/duration)；', 11, False, GRAY),
    ('仿真时长 = 字数×0.08s，后续接真实 edge-tts 只换内部实现，两端接口不变。', 11, False, GRAY),
])
footer(s, 8)

# ============================================================ 9 质量验证与实测数据
s = prs.slides.add_slide(BLANK)
title_bar(s, '质量验证与实测数据（两周累计）')
cards = [
    ('单测 47/47', 'W2 经验 8 + W3 题库/答题/决策 20\n+ W4 情感决策新增 19\n纯逻辑无 ROS 可复跑', GREEN),
    ('接口 16 项', 'W4 零新增——全部复用 W1 定义\n行为长在旧接口上\n（TTS 动作 W4 实现服务端）', BLUE),
    ('编译检查', '新增/改动 10 个 Python 文件\npy_compile 全过\n（ROS 节点亦通过）', BLUE),
    ('演示可复现', 'standalone_demo_w4.py 一条时间轴\n--fast 快速模式\nstart_demo.sh 一键含 TTS', CYAN),
]
x = 0.6
for h, body, color in cards:
    add_rect(s, x, 1.6, 2.95, 0.7, color, rounded=True, text=h, tsize=14)
    add_rect(s, x, 2.4, 2.95, 2.15, LIGHT, rounded=True)
    add_text(s, x + 0.13, 2.55, 2.69, 1.9, [(body, 10, False, NAVY)],
             align=PP_ALIGN.CENTER)
    x += 3.07
# W4 demo 终态
add_rect(s, 0.6, 4.85, 12.1, 1.0, LIGHT, rounded=True)
add_text(s, 0.9, 4.98, 11.6, 0.8, [
    ('W4 demo 终态：LV.10 ｜ 粘人型（独立→安静→粘人完整演化）｜ 情感账户 20→32（和好实测）', 13, True, NAVY),
    ('日记已落库（2025-08-28 一条）；W3 demo 终态 LV.23 / 1300 经验 / 连击 ×5 +50、×10 +100 照旧可复现。', 11.5, False, GRAY),
])
add_rect(s, 0.6, 6.05, 12.1, 0.8, ORANGE, rounded=True)
add_text(s, 0.9, 6.2, 11.5, 0.55, [
    ('⏳ 待验证：Ubuntu 侧编译 / TTS 联调 / start_demo 全流程 / 双机联调（需环境，见风险页）——两周均无新增接口，风险低。', 12, True, WHITE),
])
footer(s, 9)

# ============================================================ 10 小组成员进展与分工
s = prs.slides.add_slide(BLANK)
title_bar(s, '小组成员进展与下一步（牛浩然 / 黄鑫禹）')
# 牛浩然
add_rect(s, 0.6, 1.55, 5.9, 4.95, LIGHT, rounded=True)
add_rect(s, 0.6, 1.55, 5.9, 0.6, BLUE, rounded=True, text='牛浩然（前端 / 后端）', tsize=14)
add_text(s, 0.85, 2.3, 5.45, 4.1, [
    ('W3 已完成：', 12.5, True, NAVY),
    ('· 答题页 + 题库/判题 API（对接 /dog/quiz、/dog/combo）', 11, False, GRAY),
    ('W4 本周要干嘛（排期）：', 12.5, True, NAVY),
    ('· 日记 API：读 diary 表（W4 已落库，一天一条）→ 前端日记页', 11, False, GRAY),
    ('· 周报 API：按周汇总互动/答题/日记数据', 11, False, GRAY),
    ('· 心情时间序列 API + 前端心情曲线', 11, False, GRAY),
    ('依赖与衔接：', 12.5, True, NAVY),
    ('· diary 表结构与经验库同库（db_path），schema 已预置', 11, False, GRAY),
    ('· 决策引擎 daily_stats()/diary.latest() 可直接复用', 11, False, GRAY),
    ('· 验收：前端能展示"某天日记 + 心情曲线"即通过', 11, False, GRAY),
])
# 黄鑫禹
add_rect(s, 6.85, 1.55, 5.9, 4.95, LIGHT, rounded=True)
add_rect(s, 6.85, 1.55, 5.9, 0.6, GREEN, rounded=True, text='黄鑫禹（创意 / 测试）', tsize=14)
add_text(s, 7.1, 2.3, 5.45, 4.1, [
    ('W3 已完成：', 12.5, True, NAVY),
    ('· 语音问答链路 + 5 场景验证 + 连击体验验证', 11, False, GRAY),
    ('W4 本周要干嘛（排期）：', 12.5, True, NAVY),
    ('· 角色串联演示：表情→决策→动作→TTS 全链路走查', 11, False, GRAY),
    ('· 手感实测：抚摸/冷落/和好/仪式的节奏与话术体验', 11, False, GRAY),
    ('· 调参建议：性格轴增益/衰减、账户速率、关怀阈值', 11, False, GRAY),
    ('依赖与衔接：', 12.5, True, NAVY),
    ('· W4 已提供：性格化话术 + TTS 仿真播报 + 一键启动', 11, False, GRAY),
    ('· 验收：串联演示无卡点 + 一份调参建议清单', 11, False, GRAY),
])
add_text(s, 0.6, 6.65, 12.1, 0.5, [
    ('协作方式不变：任务清单已单独成文（docs/任务清单_W4_*.md），含依赖、验收标准、现在就能做的第一步。', 12, False, GRAY),
])
footer(s, 10)

# ============================================================ 11 风险 + 下一步 + 支持
s = prs.slides.add_slide(BLANK)
title_bar(s, '风险、下一步（W5）与需要的支持')
rows = [
    ('风险 / 待办', '影响', '应对'),
    ('Ubuntu 侧编译/联调未完成', '两周成果未在真实 ROS 环境回归', '首要待办；W4 无新增接口，预计一次编译即可验证'),
    ('性格/情感参数为初版经验值', '节奏、话术密度可能与理想手感有偏差', '黄鑫禹 W4 手感实测 → 出调参建议 → 快速迭代'),
    ('TTS 为仿真实现', '音色/语速非最终效果', '接口已定，后续替换 edge-tts 不动两端'),
    ('前端日记页依赖联调窗口', '日记 API 联调延后则页面演示顺延', '表结构与统计函数已就绪，接口面谈一次即可开工'),
]
add_table(s, rows, 0.6, 1.5, 12.1, 2.9, [3.4, 3.9, 4.8])
add_rect(s, 0.6, 4.65, 5.9, 2.1, LIGHT, rounded=True)
add_text(s, 0.85, 4.82, 5.45, 1.8, [
    ('W5 计划（对齐排期）：', 13.5, True, NAVY),
    ('① 火花计划：每日任务/打卡/奖励接入决策与结算', 11.5, False, GRAY),
    ('  （streak/daily_task/daily_reward 表 W1 已预置）', 11, False, GRAY),
    ('② 综合：两端功能联演 + 演示脚本定稿', 11.5, False, GRAY),
])
add_rect(s, 6.85, 4.65, 5.9, 2.1, ORANGE, rounded=True)
add_text(s, 7.1, 4.82, 5.45, 1.8, [
    ('需要的支持：', 13.5, True, WHITE),
    ('① Ubuntu + ROS2 环境完成编译/联调验证（两周共性待办）', 11.5, False, WHITE),
    ('② 与牛浩然/黄鑫禹各 30min 对齐接口与联调排期', 11.5, False, WHITE),
    ('③ 如需实体机演示窗口，请提前协调（W6-W7 计划内）', 11.5, False, WHITE),
])
footer(s, 11)

prs.save(OUT)

# ---------------- 边界自检：所有形状不得超出页面 ----------------
SW, SH = 13.333, 7.5
worst_r, worst_b, bad = 0.0, 0.0, []
for idx, slide in enumerate(prs.slides, 1):
    for shp in slide.shapes:
        try:
            r = (shp.left + shp.width) / 914400.0
            b = (shp.top + shp.height) / 914400.0
        except TypeError:
            continue
        worst_r, worst_b = max(worst_r, r), max(worst_b, b)
        if r > SW + 0.01 or b > SH + 0.01:
            bad.append((idx, shp.shape_type, round(r, 2), round(b, 2)))
print(f'saved: {OUT}')
print(f'bounds: max_right={worst_r:.2f}in max_bottom={worst_b:.2f}in '
      f'(page 13.33x7.50) overflows={len(bad)}')
for b in bad:
    print('  OVERFLOW:', b)
