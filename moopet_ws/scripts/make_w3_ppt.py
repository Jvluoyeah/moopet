#!/usr/bin/env python3
"""生成 W3 演示汇报 PPT（长夜月 · 李文扬）。

用法：
  python scripts/make_w3_ppt.py        # 输出 演示_W3汇报.pptx（工作区根目录）
依赖：python-pptx（ppt-master/venv 已装）
"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, '演示_W3汇报.pptx')

# 配色
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
CYAN = RGBColor(0x3B, 0xA7, 0xA5)
ORANGE = RGBColor(0xE8, 0x7A, 0x2E)
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)
GRAY = RGBColor(0x55, 0x5F, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = '微软雅黑'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _set_ea(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


def set_run(run, size, bold=False, color=NAVY, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    _set_ea(run)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, it in enumerate(lines):
        text, size, bold, color = it[0], it[1], it[2], it[3]
        sa = it[4] if len(it) > 4 else space_after
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sa)
        r = p.add_run()
        r.text = text
        set_run(r, size, bold, color)
    return tb


def add_rect(slide, x, y, w, h, fill, line=None, text='', tsize=14, tcolor=WHITE,
             bold=True, rounded=False, line_w=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_run(r, tsize, bold, tcolor)
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=BLUE, w=2.25):
    conn = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(w)
    return conn


def title_bar(slide, title, sub=''):
    add_rect(slide, 0, 0, 13.333, 0.95, NAVY, text=title, tsize=23)
    add_rect(slide, 0, 0.95, 13.333, 0.06, ORANGE)
    if sub:
        add_text(slide, 0.5, 1.12, 12.3, 0.4, [(sub, 12.5, False, GRAY)])


def footer(slide, idx):
    add_text(slide, 11.4, 7.1, 1.7, 0.3,
             [(f'{idx} / 8', 10, False, GRAY)], align=PP_ALIGN.RIGHT)


# ============================================================ 1 封面
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, 13.333, 7.5, NAVY)
add_rect(s, 0, 4.9, 13.333, 0.08, ORANGE)
add_text(s, 1.0, 1.5, 11.3, 1.1, [('W3 演示汇报：答题与基础闭环', 38, True, WHITE)])
add_text(s, 1.0, 2.55, 11.3, 0.7,
         [('感知 → 加点 → 升级 → 答题 → 决策反应', 21, True, CYAN)])
add_text(s, 1.0, 4.05, 11.3, 1.2, [
    ('长夜月 · 李文扬（架构 / 集成）', 16, False, WHITE),
    ('蓝思杯 · 赛道一：灵伴管家（智能体应用）', 14, False, WHITE),
    ('对应排期 W3（8.17–8.23）', 14, False, WHITE),
])
add_text(s, 1.0, 6.6, 11.3, 0.4,
         [('演示载体：w3/moopet_ws · standalone_demo_w3.py / start_demo.sh', 12, False, WHITE)])

# ============================================================ 2 演示定位
s = prs.slides.add_slide(BLANK)
title_bar(s, '这次演示做什么')
add_rect(s, 0.6, 1.6, 12.1, 1.25, LIGHT, rounded=True)
add_text(s, 0.9, 1.75, 11.5, 1.0, [
    ('不依赖真实机器狗硬件，用 ROS2 仿真和模拟数据，把「感知 → 加点 → 升级 → 答题 → 决策反应」', 14.5, False, NAVY),
    ('这一整条软件流程跑通，验证软件架构和养成体系已经就绪。', 14.5, False, NAVY),
])
add_rect(s, 0.6, 3.1, 12.1, 1.05, BLUE, rounded=True)
add_text(s, 0.9, 3.28, 11.5, 0.8, [
    ('思路：前两周把基础打好（接口、感知、成长），第三周把答题和决策接上，', 15, True, WHITE),
    ('让整条流程首尾连起来。', 15, True, WHITE),
])
add_text(s, 0.8, 4.5, 11.8, 1.6, [
    ('演示分三层，是同一套代码逐步加长的过程：', 14, True, NAVY),
    ('W1 接口契约 → W2 感知与成长 → W3 答题与决策', 13.5, False, GRAY),
    ('第三周的出题、判题、决策都直接复用前两周的接口和结算引擎，所以按顺序放一起演示最顺。', 13.5, False, GRAY),
])
footer(s, 2)

# ============================================================ 3 演示主线
s = prs.slides.add_slide(BLANK)
title_bar(s, '三层演示内容')
add_rect(s, 0.6, 1.5, 12.1, 4.7, WHITE, line=BLUE, line_w=1.2)
rows = [
    ('层', '周', '演示内容', '关键看点'),
    ('第一层 · 地基', 'W1', '工作空间三包 + 16 项自定义接口 + 冒烟测试', '先把接口协议定好'),
    ('第二层 · 感知与成长', 'W2', '背包 5 个感知节点 + 遛狗里程 + 8 种来源加点 + 升级/成就', '能感知，能成长'),
    ('第三层 · 答题与闭环', 'W3', '题库按等级解锁 → 答题/连击/难度降档 → 决策引擎入口 → 语音动作联动 → 状态卡', '能答题，能做反应'),
]
table_shape = s.shapes.add_table(len(rows), 4, Inches(0.8), Inches(1.65),
                                 Inches(11.7), Inches(3.3))
tbl = table_shape.table
tbl.columns[0].width = Inches(2.4)
tbl.columns[1].width = Inches(0.8)
tbl.columns[2].width = Inches(6.5)
tbl.columns[3].width = Inches(2.0)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = val
        set_run(r, 11.5 if ri == 0 else 11, ri == 0, WHITE if ri == 0 else NAVY)
add_text(s, 0.8, 5.15, 11.7, 0.7, [
    ('三层共用同一套纯逻辑（quiz_bank / exp_engine / decision），在一条时间轴上推进。', 13, True, BLUE),
])
footer(s, 3)

# ============================================================ 4 系统架构示意
s = prs.slides.add_slide(BLANK)
title_bar(s, '系统结构示意', sub='双机：背包盒子与狗本体，设相同 ROS_DOMAIN_ID 可跨机互通')

# 背包盒子
add_rect(s, 0.6, 1.9, 4.5, 3.6, CYAN, rounded=True, text='背包盒子（仿真感知）', tsize=15)
b_items = ['heartbeat · 5s 心跳', 'audio · 语音命令', 'vision · 表情/动作',
           'env · 空气质量', 'touch · 抚摸分区']
for i, t in enumerate(b_items):
    add_rect(s, 0.8, 2.62 + i * 0.6, 4.1, 0.48, LIGHT, text=t, tsize=11, tcolor=NAVY, bold=False)

# 狗本体
add_rect(s, 8.2, 1.9, 4.5, 3.6, BLUE, rounded=True, text='狗本体', tsize=15)
d_items = ['odom · 10Hz 遛狗里程', 'exp_engine · 加点/升级/成就',
           'quiz_node · 出题/判题/连击/难度', 'decision · 行为决策引擎（BT）',
           'state_mgr · 状态卡']
for i, t in enumerate(d_items):
    add_rect(s, 8.4, 2.62 + i * 0.6, 4.1, 0.48, LIGHT, text=t, tsize=11, tcolor=NAVY, bold=False)

# 中间连线：背包 → 狗
add_arrow(s, 5.15, 3.7, 8.15, 3.7, color=ORANGE, w=3.0)
add_text(s, 5.5, 3.02, 2.3, 0.4, [('感知事件', 12.5, True, GRAY)], align=PP_ALIGN.CENTER)

# 状态卡（狗本体正下方）
add_rect(s, 8.2, 6.2, 4.5, 0.75, ORANGE, rounded=True,
         text='状态卡 / 控制台（等级 / 经验 / 心情）', tsize=12)
add_arrow(s, 10.45, 5.55, 10.45, 6.15, color=GRAY, w=2.25)
add_text(s, 7.0, 5.72, 2.9, 0.35, [('/dog/state 实时刷新', 11.5, True, GRAY)],
         align=PP_ALIGN.RIGHT)
footer(s, 4)

# ============================================================ 5 闭环流程示意
s = prs.slides.add_slide(BLANK)
title_bar(s, '流程示意（演示怎么走）')
stages = [
    ('W1 地基', '16 项接口\n冒烟测试通过'),
    ('W2 感知+成长', '抚摸/散步/呼唤/握手/通风/投影/通讯\n攒经验升级，解锁题库'),
    ('W3 答题', '出题、判题\n连击 5 题 +50，10 题 +100\n连错 3 题难度降一档'),
    ('决策引擎', '进入决策引擎入口\n行为树路由\n答对表扬 / 答错鼓励'),
    ('回应与状态', 'TTS 语音\n动作联动\n状态卡刷新'),
]
x = 0.4
w = 2.23
gap = 0.35
for i, (h, body) in enumerate(stages):
    add_rect(s, x, 1.9, w, 0.72, NAVY if i != 3 else ORANGE, rounded=True, text=h, tsize=14)
    add_rect(s, x, 2.72, w, 2.5, LIGHT, rounded=True)
    add_text(s, x + 0.1, 2.88, w - 0.2, 2.2, [(body, 10, False, NAVY)],
             align=PP_ALIGN.CENTER)
    if i < 4:
        add_arrow(s, x + w + 0.04, 2.26, x + w + gap - 0.04, 2.26, color=ORANGE, w=3.0)
    x += w + gap
add_text(s, 0.4, 5.5, 12.5, 1.3, [
    ('整个流程：互动产生经验 → 升级解锁题库 → 答题连击 → 决策引擎给出回应（语音/动作）→ 状态卡刷新。', 14, True, NAVY),
    ('全程用模拟数据，不依赖真实机器狗硬件。', 13, False, GRAY),
])
footer(s, 5)

# ============================================================ 6 时间轴脚本
s = prs.slides.add_slide(BLANK)
title_bar(s, '时间轴与解说词（约 3 分钟）')
rows = [
    ('段', '时间', '环节', '屏幕效果', '解说词'),
    ('1', '0:00', '开场 + W1 地基', '三个包、16 项接口、冒烟通过', '先看地基：接口和冒烟都过了，后面的功能都建立在它上面。'),
    ('2', '0:30', 'W2 感知与成长', '背包 5 个节点；抚摸/散步等逐项加点升级', '背包盒子在模拟各种感知，经验攒起来，狗就从 LV.1 往上涨。'),
    ('3', '1:30', 'W3 答题', '等级解锁科目；连击 +50/+100；连错 3 题降档', '题库按等级开锁；连对有奖励；连错 3 题，狗会把难度调低一档陪你慢慢来。'),
    ('4', '2:20', '决策引擎 + 回应', '答题进决策引擎；行为树路由；语音动作联动', '答题不只是加分，还会进决策引擎，给出答对表扬、答错鼓励，并联动语音和动作。'),
    ('5', '2:45', '桥接 + 状态卡', '/dog/combo 往返；状态卡刷新', '顺带演示判题桥接接口（web 端也用这个），最后状态卡实时刷新。'),
]
ts = s.shapes.add_table(len(rows), 5, Inches(0.6), Inches(1.6), Inches(12.1), Inches(4.2))
tb = ts.table
for ci, wd in enumerate([0.5, 0.9, 2.0, 4.2, 4.5]):
    tb.columns[ci].width = Inches(wd)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tb.cell(ri, ci)
        cell.margin_left = Inches(0.05)
        cell.margin_right = Inches(0.05)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = val
        set_run(r, 11 if ri == 0 else 10.5, ri == 0, WHITE if ri == 0 else NAVY)
add_text(s, 0.7, 6.05, 12, 0.6, [
    ('两种跑法：独立演示 python standalone_demo_w3.py（现场/录屏，无 ROS）；ROS2 双机 start_demo.sh（看跨机互通）。', 12, True, BLUE),
])
footer(s, 6)

# ============================================================ 7 验证点与容错
s = prs.slides.add_slide(BLANK)
title_bar(s, '要验证的点与容错')
rows = [
    ('演示看点', '验证内容', '对应排期 / 目标'),
    ('题库按等级解锁', '等级门槛 3 / 8 / 15 / 20 级', '排期 W3'),
    ('连击 5 / 10 奖励', '连击与奖励逻辑', '排期 W3'),
    ('连错 3 题难度降档', '难度降档', '排期 W3'),
    ('/dog/combo 桥接', '判题接口桥接 ROS', '排期 W3'),
    ('决策引擎入口 + 行为树', '事件路由 + 语音动作联动', '排期 W3 · 李文扬'),
    ('状态实时刷新', '经验 / 状态同步延迟 ≤1s', '创意方案书指标'),
]
ts = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.5), Inches(12.1), Inches(3.4))
tb = ts.table
for ci, wd in enumerate([3.4, 4.7, 4.0]):
    tb.columns[ci].width = Inches(wd)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tb.cell(ri, ci)
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = val
        set_run(r, 11 if ri == 0 else 10.5, ri == 0, WHITE if ri == 0 else NAVY)
add_text(s, 0.6, 5.1, 12.1, 1.6, [
    ('容错：', 13.5, True, NAVY),
    ('① TTS 没就绪时只记日志，不影响主流程（TTS 节点 W4 接入）；', 12, False, GRAY),
    ('② 没有 ROS 或双机时，用独立演示顶上（同一套逻辑）；', 12, False, GRAY),
    ('③ 等级不够时答题会自动等，或演示前用预置数据；④ 单个环节出错可以跳过。', 12, False, GRAY),
])
footer(s, 7)

# ============================================================ 8 结束
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, 13.333, 7.5, NAVY)
add_rect(s, 0, 3.05, 13.333, 0.08, ORANGE)
add_text(s, 1.0, 2.0, 11.3, 1.0, [('谢谢', 38, True, WHITE)])
add_text(s, 1.0, 3.45, 11.3, 0.7, [
    ('流程已经跑通：感知 → 成长 → 答题 → 决策反应，全程模拟数据。', 17, True, CYAN)])
add_text(s, 1.0, 4.4, 11.3, 1.0, [
    ('下一步（W4）：补上决策引擎的性格、情感账户、主动关怀和纪念日，接入背包 TTS 语音问答。', 13.5, False, WHITE),
])
footer(s, 8)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print('saved:', OUT)
