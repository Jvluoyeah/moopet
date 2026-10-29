#!/usr/bin/env python3
"""生成 W3 工作汇报 PPT（李文扬 · 架构/集成，向上级汇报版）。

用法：
  python scripts/make_w3_upreport_ppt.py   # 输出  W3工作汇报_李文扬.pptx（工作区根目录）
依赖：python-pptx（ppt-master/venv 已装）
"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'W3工作汇报_李文扬.pptx')

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
CYAN = RGBColor(0x3B, 0xA7, 0xA5)
ORANGE = RGBColor(0xE8, 0x7A, 0x2E)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC0, 0x39, 0x2B)
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)
GRAY = RGBColor(0x55, 0x5F, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = '微软雅黑'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _set_ea(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


def set_run(run, size, bold=False, color=NAVY, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    _set_ea(run)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, it in enumerate(lines):
        text, size, bold, color = it[0], it[1], it[2], it[3]
        sa = it[4] if len(it) > 4 else space_after
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sa)
        r = p.add_run()
        r.text = text
        set_run(r, size, bold, color)
    return tb


def add_rect(slide, x, y, w, h, fill, line=None, text='', tsize=14, tcolor=WHITE,
             bold=True, rounded=False, line_w=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_run(r, tsize, bold, tcolor)
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=BLUE, w=2.25):
    conn = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(w)
    return conn


def title_bar(slide, title, sub=''):
    add_rect(slide, 0, 0, 13.333, 0.95, NAVY, text=title, tsize=23)
    add_rect(slide, 0, 0.95, 13.333, 0.06, ORANGE)
    if sub:
        add_text(slide, 0.5, 1.12, 12.3, 0.4, [(sub, 12.5, False, GRAY)])


def footer(slide, idx, total=9):
    add_text(slide, 11.4, 7.1, 1.7, 0.3,
             [(f'{idx} / {total}', 10, False, GRAY)], align=PP_ALIGN.RIGHT)


# ============================================================ 1 封面
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, 13.333, 7.5, NAVY)
add_rect(s, 0, 4.75, 13.333, 0.08, ORANGE)
add_text(s, 1.0, 1.35, 11.3, 1.1, [('W3 工作汇报（8.17–8.23）', 38, True, WHITE)])
add_text(s, 1.0, 2.45, 11.3, 0.8,
         [('答题 + 基础闭环 · quiz_node 集成 + 行为决策引擎骨架', 21, True, CYAN)])
add_text(s, 1.0, 3.6, 11.3, 1.0, [
    ('汇报人：李文扬（架构 / 集成）', 16, False, WHITE),
    ('长夜月 · 蓝思杯赛道一：灵伴管家（智能体应用）', 14, False, WHITE),
    ('对齐排期 W3 · 李文扬条目', 14, False, WHITE),
])
add_text(s, 1.0, 5.35, 11.3, 1.2, [
    ('本周结论：', 15, True, ORANGE),
    ('两项排期任务全部完成；「感知 → 加点 → 升级 → 答题 → 决策反应」基础闭环', 15, False, WHITE),
    ('可演示、可单测（28/28），待 Ubuntu 侧编译验证。', 15, False, WHITE),
])
add_text(s, 1.0, 6.75, 11.3, 0.4,
         [('代码：w3/moopet_ws（承接 W1 接口 + W2 感知结算层）', 12, False, WHITE)])

# ============================================================ 2 本周目标与排期对齐
s = prs.slides.add_slide(BLANK)
title_bar(s, '本周目标与排期对齐')
rows = [
    ('任务', '类型', '目标', '状态'),
    ('集成 quiz_node 并挂到决策引擎入口', '排期 W3', '答题节点接入：出题/判题/连击/难度，作答事件进决策引擎', '✅ 完成'),
    ('搭建行为决策引擎骨架（BT 目录 + 空节点）', '排期 W3', 'Behavior Tree 目录结构 + 空节点 + 决策引擎入口节点', '✅ 完成'),
    ('让闭环可跑、可演示（自加）', '补充', '题库/结算扩展/新接口/独立 demo/单测，串通整条链路', '✅ 完成'),
]
ts = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.6), Inches(12.1), Inches(2.8))
tb = ts.table
for ci, wd in enumerate([3.6, 1.1, 5.3, 2.1]):
    tb.columns[ci].width = Inches(wd)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tb.cell(ri, ci)
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = val
        set_run(r, 11.5 if ri == 0 else 11, ri == 0, WHITE if ri == 0 else NAVY)
add_rect(s, 0.6, 4.75, 12.1, 1.6, LIGHT, rounded=True)
add_text(s, 0.9, 4.95, 11.6, 1.3, [
    ('说明：', 13.5, True, NAVY),
    ('两项排期任务按计划完成；为让闭环真正可跑、可给队友演示，本周额外补齐了题库、结算扩展、', 12.5, False, GRAY),
    ('决策入口接口、无 ROS 独立演示和单测——这些都在 W1/W2 的接口与引擎上直接扩展，未另起炉灶。', 12.5, False, GRAY),
])
footer(s, 2)

# ============================================================ 3 完成情况一览
s = prs.slides.add_slide(BLANK)
title_bar(s, '完成情况一览（对照排期 / 交付物）')
rows = [
    ('交付物', '类型', '状态', '位置 / 证据'),
    ('quiz_node：出题/判题/连击/难度 + /dog/combo 桥接 + 决策事件挂入', '排期', '✅', 'src/moopet_dog/moopet_dog/quiz_node.py'),
    ('行为决策引擎骨架：BT 目录 + 空节点 + 决策引擎入口节点', '排期', '✅', 'src/moopet_dog/moopet_dog/decision/'),
    ('题库模块（7 科目 × 4 题，等级解锁/难度降档）', '补充', '✅', 'quiz_bank.py（纯逻辑）'),
    ('结算扩展：quiz 来源（+10/周上限 200）+ 连击 5/10 奖励 + 答题成就', '补充', '✅', 'exp_engine.py'),
    ('决策入口接口 DecisionEventMsg（接口总数 15→16）', '补充', '✅', 'moopet_interfaces/msg/'),
    ('无 ROS 独立演示（W1→W2→W3）+ 一键启动集成', '补充', '✅', 'standalone_demo_w3.py / start_demo.sh'),
    ('单测：题库 7 + 答题 6 + 决策 7 + 经验引擎回归 8 = 28/28', '质量', '✅', 'tests/ + boundary_test.sh'),
    ('Ubuntu 编译 / 双机 / 全流程验证', '待办', '⏳', '需 Ubuntu + ROS2 环境（见风险页）'),
]
ts = s.shapes.add_table(len(rows), 4, Inches(0.5), Inches(1.5), Inches(12.3), Inches(4.6))
tb = ts.table
for ci, wd in enumerate([5.0, 1.0, 1.0, 5.3]):
    tb.columns[ci].width = Inches(wd)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tb.cell(ri, ci)
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.01)
        cell.margin_bottom = Inches(0.01)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = val
        color = WHITE if ri == 0 else (GREEN if ri == len(rows) - 1 and val == '⏳' else NAVY)
        set_run(r, 10.5 if ri == 0 else 10, ri == 0, WHITE if ri == 0 else color)
footer(s, 3)

# ============================================================ 4 成果① 答题闭环
s = prs.slides.add_slide(BLANK)
title_bar(s, '成果① 答题闭环（quiz_node）')
add_text(s, 0.6, 1.5, 12.1, 0.5, [
    ('基于 W2 结算引擎扩展出「答题」来源，接入 决策引擎入口，形成完整闭环。', 13.5, True, BLUE)])
items = [
    ('出题 / 判题', '/dog/quiz/question、/dog/quiz/answer；判题规则在 quiz_bank（纯逻辑可单测）'),
    ('连击奖励', '连对 5 题 +50、10 题 +100（不计周上限），对齐排期'),
    ('难度降档', '连错 3 题难度降一档（3→2），带解释性反馈'),
    ('等级解锁', '语文/数学 3 级、英文/常识 8 级、科学/编程 15 级、自定义 20 级'),
    ('判题桥接', '/dog/combo 服务：web 端判题同款接口（供牛浩然答题页对接）'),
    ('挂入决策引擎', '每次作答发布 /dog/decision/event，事件进入行为决策引擎入口'),
]
for i, (h, d) in enumerate(items):
    y = 2.15 + i * 0.72
    add_rect(s, 0.6, y, 2.2, 0.6, NAVY, rounded=True, text=h, tsize=12)
    add_rect(s, 2.95, y, 9.75, 0.6, LIGHT, rounded=True)
    add_text(s, 3.1, y + 0.03, 9.5, 0.55, [(d, 11, False, NAVY)])
add_text(s, 0.6, 6.45, 12.1, 0.6, [
    ('接入点：quiz_node 复用 /dog/state 广播状态卡；加分走 exp_engine.quiz_answer（与 W2 同一引擎）。', 12, False, GRAY),
])
footer(s, 4)

# ============================================================ 5 成果② 决策引擎骨架
s = prs.slides.add_slide(BLANK)
title_bar(s, '成果② 行为决策引擎骨架（BT 目录 + 空节点）')
# 左：目录结构
add_rect(s, 0.6, 1.5, 6.0, 4.9, LIGHT, rounded=True)
add_text(s, 0.85, 1.62, 5.6, 4.7, [
    ('moopet_dog/decision/', 11.5, True, NAVY),
    ('├── behavior_tree/        # 行为树', 11, False, NAVY),
    ('│   ├── core.py           # Status/Sequence/Selector/', 11, False, NAVY),
    ('│   │                     #   Fallback/Condition/Action', 11, False, NAVY),
    ('│   ├── conditions.py     # 事件/等级/心情条件', 11, False, NAVY),
    ('│   │                     #   + W4 空节点（时间/沉默/新一天）', 11, False, NAVY),
    ('│   ├── actions.py        # 答题/抚摸/握手/呼唤/空气/散步/指引', 11, False, NAVY),
    ('│   │                     #   + W4 空节点（纪念日/主动关怀/性格）', 11, False, NAVY),
    ('│   └── trees.py          # 答题子树 + 主树', 11, False, NAVY),
    ('├── decision_engine.py    # 决策引擎入口：事件路由 → 反应', 11, False, NAVY),
    ('└── decision_engine_node.py # ROS 节点：订阅 /dog/decision/event', 11, False, NAVY),
    ('                              → TTS / /dog/action 联动', 11, False, NAVY),
])
# 右：已跑通与 W4 计划
add_rect(s, 6.95, 1.5, 5.8, 2.3, GREEN, rounded=True)
add_text(s, 7.2, 1.65, 5.3, 2.0, [
    ('本周已跑通：', 14, True, WHITE),
    ('答题事件 → 行为树路由 → 答对表扬 / 答错鼓励', 12.5, False, WHITE),
    ('→ 联动 TTS（未就绪时日志降级）与 /dog/action', 12.5, False, WHITE),
    ('→ 状态卡实时刷新', 12.5, False, WHITE),
])
add_rect(s, 6.95, 4.05, 5.8, 2.35, ORANGE, rounded=True)
add_text(s, 7.2, 4.2, 5.3, 2.1, [
    ('W4 将填充（骨架已占位）：', 14, True, WHITE),
    ('性格演化（粘人/独立/话痨/安静）', 12.5, False, WHITE),
    ('情感账户、主动关怀、纪念日触发', 12.5, False, WHITE),
    ('→ 决策输出写入成长日记', 12.5, False, WHITE),
])
footer(s, 5)

# ============================================================ 6 成果③ 可演示可验证
s = prs.slides.add_slide(BLANK)
title_bar(s, '成果③ 可演示、可验证')
cards = [
    ('16 项接口', 'W3 新增 DecisionEventMsg（决策入口）\n全部登记 CMakeLists + interface_spec'),
    ('无 ROS 独立演示', 'standalone_demo_w3.py\nW1→W2→W3 一条时间轴，Windows 可跑\n（--fast / --interactive）'),
    ('单测 28/28', '题库 7 + 答题 6 + 决策 7\n+ 经验引擎回归 8\n全部纯逻辑、无 ROS 可跑'),
    ('一键启动集成', 'start_demo.sh 新增 quiz / decision 节点\n+/dog/combo 场景，Ubuntu 上全流程'),
]
x = 0.6
w = 2.9
for i, (h, body) in enumerate(cards):
    add_rect(s, x, 1.7, w, 0.7, BLUE, rounded=True, text=h, tsize=15)
    add_rect(s, x, 2.5, w, 2.6, LIGHT, rounded=True)
    add_text(s, x + 0.15, 2.65, w - 0.3, 2.3, [(body, 10.5, False, NAVY)],
             align=PP_ALIGN.CENTER)
    x += w + 0.15
add_rect(s, 0.6, 5.5, 12.1, 1.3, LIGHT, rounded=True)
add_text(s, 0.9, 5.7, 11.6, 1.0, [
    ('演示视角：', 13.5, True, NAVY),
    ('W1 地基（接口/冒烟）→ W2 感知成长（LV.1→解锁等级）→ W3 答题/连击/难度/决策反应 → 状态卡刷新。', 12.5, False, GRAY),
    ('整条链路承接前两周成果，非另起炉灶。', 12.5, False, GRAY),
])
footer(s, 6)

# ============================================================ 7 质量验证与实测数据
s = prs.slides.add_slide(BLANK)
title_bar(s, '质量验证与实测数据')
add_rect(s, 0.6, 1.5, 6.0, 2.4, GREEN, rounded=True)
add_text(s, 0.85, 1.65, 5.5, 2.1, [
    ('本地验证（Windows，纯逻辑）：', 13.5, True, WHITE),
    ('· 单测 28/28 全过（含 W2 经验引擎回归 8/8）', 12, False, WHITE),
    ('· 全部新增 Python 文件 py_compile 通过', 12, False, WHITE),
    ('· 题库模块自带自检通过', 12, False, WHITE),
])
add_rect(s, 6.95, 1.5, 5.8, 2.4, LIGHT, rounded=True)
add_text(s, 7.2, 1.65, 5.35, 2.1, [
    ('demo 实测（--fast 一键复现）：', 13.5, True, NAVY),
    ('· 连击×5 → +50；连击×10 → +100', 12, False, GRAY),
    ('· 连错 3 题 → 难度 3→2', 12, False, GRAY),
    ('· 中途升级解锁 科学/编程/自定义', 12, False, GRAY),
])
add_rect(s, 0.6, 4.15, 12.1, 1.5, LIGHT, rounded=True)
add_text(s, 0.9, 4.35, 11.5, 1.2, [
    ('最终状态卡：LV.23 / 累计 1300 经验 / 本周 380/2000', 14, True, NAVY),
    ('成就解锁 4 项（全能互动 / 经验丰富 / 答题小能手 / 大师养成）；/dog/combo 桥接往返正常。', 12.5, False, GRAY),
])
add_rect(s, 0.6, 5.9, 12.1, 0.9, ORANGE, rounded=True)
add_text(s, 0.9, 6.05, 11.5, 0.6, [
    ('⏳ 待验证：Ubuntu 侧 colcon build（16 接口）/ 冒烟 / start_demo 全流程 / 双机联调（需环境，见风险页）。', 12.5, True, WHITE),
])
footer(s, 7)

# ============================================================ 8 风险与待办
s = prs.slides.add_slide(BLANK)
title_bar(s, '风险与待办')
rows = [
    ('风险 / 待办', '影响', '应对 / 状态'),
    ('Ubuntu 侧编译与实机验证未完成', '新增接口/节点未在真实 ROS 环境验证', '首要待办；需 Ubuntu+ROS2 环境，预计下周初完成'),
    ('背包 TTS 节点未实现', '决策→语音联动暂为日志降级', 'W4 接入；不影响答题/动作闭环主线'),
    ('并行协作接口对接', '牛浩然答题页、黄鑫禹语音 demo 依赖新接口', '接口已就绪并登记规范，已发任务清单'),
    ('双机联调需两台设备', '双机跨机互通未实测', 'W6 双机联调；单机仿真可先行验证'),
]
ts = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.6), Inches(12.1), Inches(3.6))
tb = ts.table
for ci, wd in enumerate([4.2, 3.6, 4.3]):
    tb.columns[ci].width = Inches(wd)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tb.cell(ri, ci)
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = val
        set_run(r, 11 if ri == 0 else 10.5, ri == 0, WHITE if ri == 0 else NAVY)
add_text(s, 0.6, 5.5, 12.1, 0.9, [
    ('结论：本周排期目标全部达成，风险可控；唯一硬性依赖是 Ubuntu 环境验证。', 13.5, True, RED),
])
footer(s, 8)

# ============================================================ 9 下一步与需要的支持
s = prs.slides.add_slide(BLANK)
title_bar(s, '下一步（W4）与需要的支持')
add_rect(s, 0.6, 1.5, 12.1, 3.1, LIGHT, rounded=True)
add_text(s, 0.9, 1.7, 11.5, 2.8, [
    ('W4 计划（对齐排期）：', 15, True, NAVY),
    ('① 决策引擎完稿：性格（粘人/独立/话痨/安静）由计数器+衰减演化；', 13, False, GRAY),
    ('② 情感账户机制（依赖感/平衡逻辑）；', 13, False, GRAY),
    ('③ 主动关怀（长沉默/特殊场景）、纪念日触发（早安/晚安/里程碑/自定义）；', 13, False, GRAY),
    ('④ 接入背包 TTS，双端联调：决策输出 → /dog/action + TTS。', 13, False, GRAY),
])
add_rect(s, 0.6, 4.85, 12.1, 1.9, ORANGE, rounded=True)
add_text(s, 0.9, 5.05, 11.5, 1.6, [
    ('需要的支持：', 15, True, WHITE),
    ('① Ubuntu + ROS2 环境（或授权使用开发机）完成编译/冒烟/全流程验证；', 13, False, WHITE),
    ('② 与牛浩然（日记/周报 API）、黄鑫禹（串联演示）对齐接口与联调时间。', 13, False, WHITE),
])
footer(s, 9)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print('saved:', OUT)
