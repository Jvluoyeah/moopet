#!/usr/bin/env python3
"""生成 W5 工作汇报 PPT（李文扬 · 架构/集成，向上级汇报版）。

背景：W3+W4 已合并汇报过；本周汇报 W5（8.31–9.6 火花 + 综合）。
核心：火花计划引擎（连续陪伴/三档提醒/每日任务/签到抽奖 200 保底/里程碑解锁）
      + 双栈集成（决策引擎 + 火花 + 纪念日统一调度）+ 抽奖经验不计周上限。

用法：
  python scripts/make_w5_upreport_ppt.py   # 输出 W5工作汇报_李文扬.pptx
依赖：python-pptx（已装）
"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'W5工作汇报_李文扬.pptx')

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
CYAN = RGBColor(0x3B, 0xA7, 0xA5)
ORANGE = RGBColor(0xE8, 0x7A, 0x2E)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC0, 0x39, 0x2B)
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)
GRAY = RGBColor(0x55, 0x5F, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = '微软雅黑'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 11


def _set_ea(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


def set_run(run, size, bold=False, color=NAVY, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    _set_ea(run)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, it in enumerate(lines):
        text, size, bold, color = it[0], it[1], it[2], it[3]
        sa = it[4] if len(it) > 4 else space_after
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sa)
        r = p.add_run()
        r.text = text
        set_run(r, size, bold, color)
    return tb


def add_rect(slide, x, y, w, h, fill, line=None, text='', tsize=14, tcolor=WHITE,
             bold=True, rounded=False, line_w=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_run(r, tsize, bold, tcolor)
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=BLUE, w=2.25):
    conn = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(w)
    return conn


def title_bar(slide, title, sub=''):
    add_rect(slide, 0, 0, 13.333, 0.95, NAVY, text=title, tsize=23)
    add_rect(slide, 0, 0.95, 13.333, 0.06, ORANGE)
    if sub:
        add_text(slide, 0.5, 1.12, 12.3, 0.4, [(sub, 12.5, False, GRAY)])


def footer(slide, idx):
    add_text(slide, 11.4, 7.1, 1.7, 0.3,
             [(f'{idx} / {TOTAL}', 10, False, GRAY)], align=PP_ALIGN.RIGHT)


def add_table(slide, rows, x, y, w, h, col_widths, header_size=11.5, body_size=10.5,
              row_h=None):
    ts = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                Inches(w), Inches(h))
    tb = ts.table
    for ci, wd in enumerate(col_widths):
        tb.columns[ci].width = Inches(wd)
    for ri, row in enumerate(rows):
        if row_h and ri == 0:
            tb.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = tb.cell(ri, ci)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = val
            set_run(r, header_size if ri == 0 else body_size, ri == 0,
                    WHITE if ri == 0 else NAVY)
    return tb


# ============================================================ 1 封面
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, 13.333, 7.5, NAVY)
add_rect(s, 0, 4.75, 13.333, 0.08, ORANGE)
add_text(s, 1.0, 1.2, 11.3, 1.1, [('W5 工作汇报（8.31–9.6）', 38, True, WHITE)])
add_text(s, 1.0, 2.35, 11.3, 0.8,
         [('火花计划 + 综合 · 连续陪伴 · 三档提醒 · 签到抽奖 · 里程碑解锁', 21, True, CYAN)])
add_text(s, 1.0, 3.5, 11.3, 1.15, [
    ('汇报人：李文扬（架构 / 集成）', 16, False, WHITE),
    ('长夜月 · 蓝思杯赛道一：灵伴管家（智能体应用）', 14, False, WHITE),
    ('对齐排期 W5 · 李文扬条目（W3+W4 已合并汇报，本次为 W5 单独汇报）', 13, False, WHITE),
])
add_text(s, 1.0, 5.3, 11.3, 1.2, [
    ('单周结论：', 15, True, ORANGE),
    ('W5 火花计划按排期完成——从"会陪伴"到"记得陪你、催你回来续火"，', 15, False, WHITE),
    ('并完成决策引擎 + 火花 + 纪念日双栈统一调度（75/75 单测全过）。', 15, False, WHITE),
])
add_text(s, 1.0, 6.85, 11.3, 0.4,
         [('代码：w5/moopet_ws（周周快照，W1-W4 成果全部保留）', 12, False, WHITE)])

# ============================================================ 2 单周结论（一页看懂）
s = prs.slides.add_slide(BLANK)
title_bar(s, '单周结论（一页看懂）')
# 左卡：火花计划
add_rect(s, 0.6, 1.6, 5.9, 3.0, LIGHT, rounded=True)
add_rect(s, 0.6, 1.6, 5.9, 0.62, ORANGE, rounded=True, text='火花计划引擎（SparkEngine）', tsize=14)
add_text(s, 0.85, 2.4, 5.45, 2.1, [
    ('· 连续陪伴：点燃 / 续火 / 断档熄灭 / 回归挽回', 12, False, NAVY),
    ('· 0~5 档佩戴：连续天数//7，封顶 5 档', 12, False, NAVY),
    ('· 三档提醒：18:00 / 21:00 / 23:00 催续火', 12, False, NAVY),
    ('· 每日任务：4 点随机刷新 2 个，完成 +2 抽奖机会', 12, False, NAVY),
    ('· 签到抽奖：概率出奖 + 200 次保底超级大奖', 12, False, NAVY),
])
# 右卡：双栈集成
add_rect(s, 6.85, 1.6, 5.9, 3.0, LIGHT, rounded=True)
add_rect(s, 6.85, 1.6, 5.9, 0.62, GREEN, rounded=True, text='双栈集成（决策引擎 + 火花 + 纪念日）', tsize=14)
add_text(s, 7.1, 2.4, 5.45, 2.1, [
    ('· build_spark_tree 并入决策引擎主树（置前）', 12, False, NAVY),
    ('· on_tick 统一调度：火花提醒 + 纪念日 + 主动关怀', 12, False, NAVY),
    ('· 互动事件 → 连续陪伴 + 每日任务进度双向驱动', 12, False, NAVY),
    ('· 抽奖经验 weekly-exempt：不计每周上限', 12, False, NAVY),
    ('· 里程碑解锁：功能 / LED 灯位 / 肤位占位', 12, False, NAVY),
])
# 底部结论条
add_rect(s, 0.6, 4.9, 12.1, 1.75, NAVY, rounded=True)
add_text(s, 0.95, 5.12, 11.5, 1.4, [
    ('一周合起来看：', 14.5, True, ORANGE),
    ('W4 让狗"懂感情"，W5 让狗"记得陪伴、会提醒、有成长外观"——排期 9.18 前的软件主线', 13, False, WHITE),
    ('（感知→成长→答题→情感→火花陪伴）已全部打通，剩 W6 稳定性收尾 + W7 打磨提交。', 13, False, WHITE),
])
footer(s, 2)

# ============================================================ 3 排期对照
s = prs.slides.add_slide(BLANK)
title_bar(s, '排期对照（W5 · 李文扬条目）')
rows = [
    ('周次', '排期任务', '完成情况', '状态'),
    ('W5', '火花全流程与行为/回复联动（提醒三档 18/21/23）', '连续陪伴点燃/续火/熄灭/挽回 + 档位；18/21/23 命中且当天未续火才提醒、同档去重', '✅'),
    ('W5', '双栈集成：决策引擎 + 火花 + 纪念日统一调度', 'build_spark_tree 并入主树并置前；on_tick 统一检查；互动事件驱动连续陪伴+每日任务', '✅'),
    ('W5', '里程碑解锁玩法接入（LED 立场/肤位数据占位）', '核心前 10 级 + 每 5 级一项到 30；LED 灯位 6 处 + 肤位 7 色数据占位', '✅'),
    ('附加', '签到 / 抽奖（狗2.md：每日任务 +2、签到 +1，200 次保底）', 'draw_lottery 概率 40/30/20/9.5/0.5%；spark_meta.pity 保底计数；经验不计周上限', '✅'),
    ('附加', '抽奖经验 weekly-exempt', 'exp_engine 新增 WEEKLY_EXEMPT_SOURCES + lottery_gain()', '✅'),
]
add_table(s, rows, 0.6, 1.55, 12.1, 4.3, [0.9, 4.2, 5.6, 1.4])
add_text(s, 0.6, 6.2, 12.1, 0.6, [
    ('说明：W5 三项排期任务全部完成，无延期；附加完成签到抽奖（streak/daily_task/daily_reward 表 W1 已预置，直接复用）。', 12.5, False, GRAY),
])
footer(s, 3)

# ============================================================ 4 W5 成果① 火花计划引擎
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W5 成果① 火花计划引擎：连续陪伴 + 三档提醒（对齐排期「提醒三档 18/21/23」）')
# 左：连续陪伴状态机
add_rect(s, 0.6, 1.6, 5.9, 4.6, LIGHT, rounded=True)
add_text(s, 0.85, 1.75, 5.4, 0.4, [('连续陪伴状态机（register_activity / settle_missed_days）', 13, True, NAVY)])
add_rect(s, 0.85, 2.3, 2.4, 1.0, GREEN, rounded=True, text='首次点燃\nnew（连续 1 天）', tsize=10.5)
add_rect(s, 3.5, 2.3, 2.7, 1.0, BLUE, rounded=True, text='续火 keep\n连续互动 +1', tsize=10.5)
add_arrow(s, 3.25, 2.8, 3.5, 2.8, color=GREEN, w=2.0)
add_rect(s, 0.85, 3.7, 2.4, 1.0, RED, rounded=True, text='熄灭 extinguish\n断档 ≥2 天归零', tsize=10.5)
add_arrow(s, 2.2, 3.3, 2.05, 3.7, color=RED, w=2.0)
add_rect(s, 3.5, 3.7, 2.7, 1.0, ORANGE, rounded=True, text='挽回 recover\n回归互动重燃', tsize=10.5)
add_arrow(s, 4.0, 3.3, 4.4, 3.7, color=ORANGE, w=2.0)
add_text(s, 0.85, 4.95, 5.45, 1.2, [
    ('· 档位 = 连续天数//7，封顶 5 档（0~5 佩戴）', 11, False, GRAY),
    ('· 历史最高天数保留（max_days），挽回不丢纪录', 11, False, GRAY),
    ('· 数据落 streak 表（W1 预置），与 exp_engine 同口径', 11, False, GRAY),
])
# 右：三档提醒
add_rect(s, 6.85, 1.6, 5.9, 4.6, LIGHT, rounded=True)
add_text(s, 7.1, 1.75, 5.4, 0.4, [('三档提醒 18:00 / 21:00 / 23:00（当天未续火才提醒）', 13, True, NAVY)])
add_rect(s, 7.1, 2.3, 5.35, 0.95, CYAN, rounded=True,
         text='18:00  温和：「主人，今天还没陪我呢，来点个火续上连续陪伴吧～」', tsize=10.5)
add_rect(s, 7.1, 3.4, 5.35, 0.95, BLUE, rounded=True,
         text='21:00  变弱：「主人，都这个点了，再不来续火，我们的火花就要变弱啦。」', tsize=10.5)
add_rect(s, 7.1, 4.5, 5.35, 0.95, RED, rounded=True,
         text='23:00  最后：「主人，最后一小时啦！再不续火，今天的火花就要熄灭咯。」', tsize=10.5)
add_text(s, 7.1, 5.6, 5.45, 0.6, [
    ('去重：同档当天只提醒一次（mark_reminded）；已互动当天不提醒。', 10.5, False, GRAY),
])
footer(s, 4)

# ============================================================ 5 W5 成果② 每日任务 + 签到抽奖
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W5 成果② 每日任务 + 签到抽奖（对齐狗2.md 抽奖规则）')
# 左：每日任务
add_rect(s, 0.6, 1.6, 5.9, 4.6, LIGHT, rounded=True)
add_rect(s, 0.6, 1.6, 5.9, 0.62, GREEN, rounded=True, text='每日任务（每天 4 点随机刷新 2 个）', tsize=14)
add_text(s, 0.85, 2.4, 5.45, 3.7, [
    ('任务池：抚摸 / 握手 / 呼唤 / 散步 / 答题 / 通风', 11.5, False, NAVY),
    ('· report_event(source) 按互动事件累计进度', 11.5, False, NAVY),
    ('· 答对才算 quiz、通风完成才算 ventilate', 11.5, False, NAVY),
    ('· 全部完成 → +2 次抽奖机会（task_bonus 只发一次）', 11.5, False, NAVY),
    ('数据：daily_task 表（W1 预置，UNIQUE(date, task_type)）', 11, False, GRAY),
])
# 右：签到抽奖
add_rect(s, 6.85, 1.6, 5.9, 4.6, LIGHT, rounded=True)
add_rect(s, 6.85, 1.6, 5.9, 0.62, ORANGE, rounded=True, text='签到 + 抽奖（40/30/20/9.5/0.5% · 200 次保底）', tsize=14)
add_text(s, 7.1, 2.4, 5.45, 1.5, [
    ('· 每日签到 → +1 次抽奖机会（sign_in 每天 1 次）', 11.5, False, NAVY),
    ('· 完成每日养成 → +2 次（与任务加成合并）', 11.5, False, NAVY),
])
add_rect(s, 7.1, 4.0, 5.35, 1.55, LIGHT, rounded=True)
add_text(s, 7.3, 4.12, 5.0, 1.35, [
    ('奖池：三等奖+20(40%) / 二等奖+50(30%) / 一等奖+100(20%)', 10.5, False, GRAY),
    ('特等奖+500(9.5%) / 超级大奖+10000(0.5%)', 10.5, False, GRAY),
    ('200 次未中 → 第 200 次保底出超级大奖（spark_meta.pity）', 10.5, False, NAVY),
    ('抽奖经验 weekly-exempt：加累计经验/等级，不计周上限', 10.5, False, NAVY),
])
footer(s, 5)

# ============================================================ 6 W5 成果③ 里程碑解锁玩法
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W5 成果③ 里程碑解锁玩法接入（LED 灯位 / 肤位数据占位）')
add_rect(s, 0.6, 1.55, 12.1, 1.15, NAVY, rounded=True)
add_text(s, 0.9, 1.7, 11.5, 0.9, [
    ('对齐狗2.md 等级解锁：核心功能 1~9 级解锁、10 级核心完成；后续非核心每 5 级一项到 30 级。', 12.5, True, WHITE),
    ('milestone_unlocks(level) 纯函数返回：功能列表 + LED 灯位 + 肤位外观（数据占位，供前端消费）。', 11.5, False, WHITE),
])
# 等级表
rows = [
    ('等级', '解锁玩法（功能）', 'LED 灯位', '肤位外观'),
    ('LV.1', '基础功能（时钟/天气/智能家居管理）', '耳部 LED', '经典白'),
    ('LV.5', '+监控/空气质量检测…（7 项）', '眼部 LED', '暖橙'),
    ('LV.10', '核心功能解锁完成（12 项）', '胸背 LED', '星空蓝'),
    ('LV.15', '非核心功能·A', '四肢 LED', '薄荷绿'),
    ('LV.20', '非核心功能·B', '尾部 LED', '曜石黑'),
    ('LV.30', '非核心功能·D（到 30 级完整）', '全身 LED', '鎏金'),
]
add_table(s, rows, 0.6, 2.95, 12.1, 3.0, [1.2, 5.2, 2.8, 2.9], header_size=11, body_size=10)
add_rect(s, 0.6, 6.15, 12.1, 0.9, LIGHT, rounded=True)
add_text(s, 0.9, 6.3, 11.5, 0.65, [
    ('定位：LED 灯位 / 肤位为数据占位——前端主题/皮肤切换与"机器狗外观随等级成长"先按占位画槽位/色块，最终资源由设计侧落地。', 11.5, False, GRAY),
])
footer(s, 6)

# ============================================================ 7 W5 成果④ 双栈集成
s = prs.slides.add_slide(BLANK)
title_bar(s, 'W5 成果④ 双栈集成：决策引擎 + 火花 + 纪念日统一调度')
# 链路图：决策引擎主树
nodes = [
    (0.6, 'on_tick 周期检查\n（节点 60s 定时器）', BLUE),
    (3.35, '行为决策引擎主树\n（BT Selector）', GREEN),
    (6.35, '火花子树\nbuild_spark_tree\n三档提醒 + 状态联动', ORANGE),
    (9.35, '纪念日/仪式/关怀\n（W4 已有）', CYAN),
]
y1 = 2.0
for x, txt, color in nodes:
    add_rect(s, x, y1, 2.5, 1.5, color, rounded=True, text=txt, tsize=11)
add_arrow(s, 3.1, y1 + 0.75, 3.35, y1 + 0.75, w=2.5)
add_arrow(s, 5.85, y1 + 0.75, 6.35, y1 + 0.75, w=2.5)
add_arrow(s, 8.85, y1 + 0.75, 9.35, y1 + 0.75, w=2.5)
# 说明
add_rect(s, 0.6, 4.0, 12.1, 1.3, LIGHT, rounded=True)
add_text(s, 0.9, 4.15, 11.5, 1.05, [
    ('行为口径（重要）：', 12.5, True, NAVY),
    ('火花提醒子树置前——今天还没互动时，18/21/23 的催续火优先于晚安仪式/深夜关怀', 11.5, False, GRAY),
    ('（"没回家的人先喊回来续火，不道晚安"）；已互动则火花不触发，照常走仪式。', 11.5, False, GRAY),
])
# 事件驱动
add_rect(s, 0.6, 5.5, 12.1, 1.1, NAVY, rounded=True)
add_text(s, 0.9, 5.65, 11.5, 0.85, [
    ('互动事件双向驱动：', 12.5, True, ORANGE),
    ('决策事件 → register_activity（连续陪伴）+ report_event（每日任务进度）→ 火花状态事件 → 主树联动话术（续火/点燃/挽回/熄灭）。', 11.5, False, WHITE),
])
footer(s, 7)

# ============================================================ 8 质量验证与实测数据
s = prs.slides.add_slide(BLANK)
title_bar(s, '质量验证与实测数据（W5 累计）')
cards = [
    ('单测 75/75', 'W5 新增 28（火花 18\n+ 决策集成 10）\n旧 47 项不回归\n纯逻辑无 ROS 可复跑', GREEN),
    ('接口 0 新增', '全部复用 W1 接口与表\n（streak/daily_task/\ndaily_reward + 自建\nspark_meta 兜底）', BLUE),
    ('编译检查', '新增/改动 11 个\nPython 文件\npy_compile 全过\n（ROS 节点亦通过）', BLUE),
    ('演示可复现', 'standalone_demo_w5.py\n六幕时间轴\n--fast 快速模式\nstart_demo.sh 一键', CYAN),
]
x = 0.6
for h, body, color in cards:
    add_rect(s, x, 1.6, 2.95, 0.7, color, rounded=True, text=h, tsize=14)
    add_rect(s, x, 2.4, 2.95, 2.15, LIGHT, rounded=True)
    add_text(s, x + 0.13, 2.55, 2.69, 1.9, [(body, 10, False, NAVY)],
             align=PP_ALIGN.CENTER)
    x += 3.07
# W5 demo 终态
add_rect(s, 0.6, 4.85, 12.1, 1.0, LIGHT, rounded=True)
add_text(s, 0.9, 4.98, 11.6, 0.8, [
    ('W5 demo 实测：点燃→续火→升档 → 三档提醒（18/21/23 依次催续火）→ 断档熄灭→挽回重燃', 12.5, True, NAVY),
    ('→ 每日任务完成 +2 机会 → 签到+抽奖（三等/二等/特等实时加点升级，周经验不变）→ LV.20 解锁 14 功能 + 5 LED + 5 肤位。', 11, False, GRAY),
])
add_rect(s, 0.6, 6.05, 12.1, 0.8, ORANGE, rounded=True)
add_text(s, 0.9, 6.2, 11.5, 0.55, [
    ('⏳ 待验证：Ubuntu 侧编译 / 决策节点火花提醒日志 / start_demo 全流程 / 双机联调（需环境，见风险页）——W5 无新增接口，风险低。', 12, True, WHITE),
])
footer(s, 8)

# ============================================================ 9 小组成员进展与分工
s = prs.slides.add_slide(BLANK)
title_bar(s, '小组成员进展与下一步（牛浩然 / 黄鑫禹）')
# 牛浩然
add_rect(s, 0.6, 1.55, 5.9, 4.95, LIGHT, rounded=True)
add_rect(s, 0.6, 1.55, 5.9, 0.6, BLUE, rounded=True, text='牛浩然（前端 / 后端）', tsize=14)
add_text(s, 0.85, 2.3, 5.45, 4.1, [
    ('W5 本周要干嘛（排期）：', 12.5, True, NAVY),
    ('· 后端：每日任务 API（4 点刷新/随机任务）+ 成就查询 API', 11, False, GRAY),
    ('· 后端：签到/抽奖 API（每日签到 1、任务完成 +2，200 保底）', 11, False, GRAY),
    ('· 后端：火花 API（连续天数/0-5 档/续火/熄灭/挽回）', 11, False, GRAY),
    ('· 前端：火花展示 + 每日任务面板 + 成就徽章页', 11, False, GRAY),
    ('· 前端：主题/皮肤切换 + 外观随等级成长', 11, False, GRAY),
    ('依赖与衔接：', 12.5, True, NAVY),
    ('· SparkEngine 已就绪（18 项单测锁定），直接包一层 API', 11, False, GRAY),
    ('· lottery_gain() 已加好，抽奖经验不计周上限', 11, False, GRAY),
    ('· LED/肤位先用占位画槽位，最终资源待设计侧补', 11, False, GRAY),
])
# 黄鑫禹
add_rect(s, 6.85, 1.55, 5.9, 4.95, LIGHT, rounded=True)
add_rect(s, 6.85, 1.55, 5.9, 0.6, GREEN, rounded=True, text='黄鑫禹（创意 / 测试）', tsize=14)
add_text(s, 7.1, 2.3, 5.45, 4.1, [
    ('W5 本周要干嘛（排期）：', 12.5, True, NAVY),
    ('· 端到端回归：遛狗→抚摸→答题→升级→火花 一整套回放', 11, False, GRAY),
    ('· 边界：断网、跨日、跨周、重复触发防重复加点', 11, False, GRAY),
    ('依赖与衔接：', 12.5, True, NAVY),
    ('· W5 已提供：火花全链路 + 六幕 demo + 一键启动', 11, False, GRAY),
    ('· 跨日/跨周/防重复可直接用 spark 的 date 参数单测', 11, False, GRAY),
    ('手感反馈：', 12.5, True, NAVY),
    ('· 三档提醒一天催三次的密度、断档一天就熄灭的宽容度', 11, False, GRAY),
    ('· 抽奖概率节奏（可脚本模拟 N 次看分布）→ 调参建议', 11, False, GRAY),
])
add_text(s, 0.6, 6.65, 12.1, 0.5, [
    ('协作方式不变：任务照会已单独成文（docs/任务清单_W5_*.md），含依赖、验收标准、现在就能做的第一步。', 12, False, GRAY),
])
footer(s, 9)

# ============================================================ 10 风险 + 下一步 + 支持
s = prs.slides.add_slide(BLANK)
title_bar(s, '风险、下一步（W6）与需要的支持')
rows = [
    ('风险 / 待办', '影响', '应对'),
    ('Ubuntu 侧编译/联调未完成', 'W5 成果未在真实 ROS 环境回归', '首要待办；W5 无新增接口，预计一次编译即可验证'),
    ('火花/抽奖参数为初版经验值', '提醒密度、保底节奏可能与理想手感有偏差', '黄鑫禹 W5 手感实测 → 出调参建议 → 快速迭代'),
    ('LED/肤位为数据占位', '外观效果非最终资源', '前端先按占位画槽位，最终资源设计侧补，不阻塞联调'),
    ('每日任务池固定 6 种', '玩法多样性有限', '任务池可扩展，W6 视体验反馈补充'),
]
add_table(s, rows, 0.6, 1.5, 12.1, 2.9, [3.4, 3.9, 4.8])
add_rect(s, 0.6, 4.65, 5.9, 2.1, LIGHT, rounded=True)
add_text(s, 0.85, 4.82, 5.45, 1.8, [
    ('W6 计划（对齐排期）：', 13.5, True, NAVY),
    ('① 双机真实局域网联通、心跳保活断连重连', 11.5, False, GRAY),
    ('② 排行榜 + WebSocket 重连 + 抽奖概率/刷级可视化验证', 11.5, False, GRAY),
    ('③ 完整演示脚本 + 预置演示数据（直升等级/预置火花）', 11.5, False, GRAY),
])
add_rect(s, 6.85, 4.65, 5.9, 2.1, ORANGE, rounded=True)
add_text(s, 7.1, 4.82, 5.45, 1.8, [
    ('需要的支持：', 13.5, True, WHITE),
    ('① Ubuntu + ROS2 环境完成编译/联调验证（共性待办）', 11.5, False, WHITE),
    ('② 与牛浩然/黄鑫禹各 30min 对齐火花 API 与联调排期', 11.5, False, WHITE),
    ('③ 如需实体机演示窗口，请提前协调（W6-W7 计划内）', 11.5, False, WHITE),
])
footer(s, 10)

# ============================================================ 11 结尾（谢谢 / 联系）
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, 13.333, 7.5, NAVY)
add_text(s, 1.0, 2.4, 11.3, 1.0, [('谢谢 · 请多指教', 40, True, WHITE)])
add_text(s, 1.0, 3.7, 11.3, 1.0, [
    ('W5 一句话：从"会陪伴"到"记得陪你、催你续火、陪你长大"。', 17, True, CYAN),
    ('代码 w5/moopet_ws · 单测 75/75 · demo standalone_demo_w5.py', 14, False, WHITE),
])
add_text(s, 1.0, 5.6, 11.3, 0.6, [
    ('汇报人：李文扬（架构 / 集成）· 长夜月 · 蓝思杯赛道一：灵伴管家', 13, False, WHITE),
])
add_text(s, 1.0, 6.4, 11.3, 0.4,
         [('承接 w4/moopet_ws，W1-W4 成果全部保留（周周快照）。', 11, False, WHITE)])

prs.save(OUT)

# ---------------- 边界自检：所有形状不得超出页面 ----------------
SW, SH = 13.333, 7.5
worst_r, worst_b, bad = 0.0, 0.0, []
for idx, slide in enumerate(prs.slides, 1):
    for shp in slide.shapes:
        try:
            r = (shp.left + shp.width) / 914400.0
            b = (shp.top + shp.height) / 914400.0
        except TypeError:
            continue
        worst_r, worst_b = max(worst_r, r), max(worst_b, b)
        if r > SW + 0.01 or b > SH + 0.01:
            bad.append((idx, shp.shape_type, round(r, 2), round(b, 2)))
print(f'saved: {OUT}')
print(f'bounds: max_right={worst_r:.2f}in max_bottom={worst_b:.2f}in '
      f'(page 13.33x7.50) overflows={len(bad)}')
for b in bad:
    print('  OVERFLOW:', b)
